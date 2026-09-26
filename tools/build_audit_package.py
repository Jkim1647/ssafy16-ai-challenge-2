"""Day 1 — 라벨 감사 500개 패키지 생성: 구글시트용 CSV 1개 + 감사 대상 이미지만 담은 zip 1개.

**Kaggle 기본규칙 4-b: validation·test 레코드의 수동 라벨링 금지.** 사람이 보는 것은 train·dev뿐이며,
test가 한 행이라도 섞이면 `_forbid_disallowed_split()`이 에러로 멈춘다(tools/build_label_sheet.py와 동일).

표본 구성(총 500, seed 고정이라 다시 돌려도 같은 500개가 나온다):
  1. baseline300  — dev에서 day1_baseline_9b와 같은 함수·seed로 뽑은 300개(유형 비율 유지).
                    9B baseline의 accuracy 계산 표본과 동일해야 label noise 상한이 그 표본에 그대로 대응한다.
  2. supplement   — baseline300 안에서 10개 미만인 희소 유형을 dev에서 보충(감사 전용).
  3. forced       — 압축률 이상치·train/dev 중복 이미지·크기 극단(감사 전용, forced_reason 표시).
  4. train_random — 위 세 가지를 합친 뒤 500까지 train에서 무작위로 채운다. train 라벨 노이즈를 dev와
                    따로 추정하기 위한 표본이다.

label noise 추정치는 **baseline300 + train_random**(무작위 표본)으로만 계산한다. supplement·forced는
의도적으로 편향된 표본이라 noise 비율에 넣으면 상한이 왜곡된다(tools/merge_label_sheets.py가 분리한다).

담당 배정: 건순·민우·현호 3명에게 라운드로빈으로 각 ~167개. 같은 출처(source)가 세 사람에게 고르게 간다.

산출물 (data_meta/label_audit/):
  - label_audit_500.xlsx         작성용 엑셀(드롭다운·필터 포함). 구글 드라이브에 올려 구글시트로 열면 된다
  - label_audit_500.csv          같은 내용의 CSV(드롭다운 없음, 가져오기용 예비)
  - label_audit_images_500.zip   이 500장만 든 zip. 기본은 AES 암호 zip (암호는 실행 시 입력, MM DM으로 공유)

실행:
    python tools/build_audit_package.py               # 암호 zip
    python tools/build_audit_package.py --no-password  # 암호 없이(팀 드라이브 내부 전달일 때만)

선행: 스키마 확정(configs/base.yaml data.columns), tools/classify_questions.py 실행,
      (권장) tools/extract_duplicates.py 실행.
"""

from __future__ import annotations

import argparse
import csv
import getpass
import hashlib
import sys
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any

TOOLS_DIR = Path(__file__).resolve().parent
REPO_ROOT = TOOLS_DIR.parent
sys.path.insert(0, str(REPO_ROOT))
sys.path.insert(0, str(TOOLS_DIR))

for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(encoding="utf-8")  # Windows 콘솔(cp949)에서 한글 깨짐 방지

import build_label_sheet as bls  # noqa: E402

TARGET_N = 500
ANNOTATORS = ["건순", "민우", "현호"]
OUT_DIR = REPO_ROOT / "data_meta" / "label_audit"
CSV_PATH = OUT_DIR / "label_audit_500.csv"
XLSX_PATH = OUT_DIR / "label_audit_500.xlsx"
ZIP_PATH = OUT_DIR / "label_audit_images_500.zip"

SOURCE_ORDER = ["baseline300", "supplement", "forced", "train_random"]

FIELDS = [
    "번호", "row_id", "split", "image_file", "질문",
    "선택지1", "선택지2", "선택지3", "선택지4", "제공된정답",
    "auto_type", "matched_rule", "source", "forced_reason",
    "담당", "유형확정", "라벨정확", "판독가능", "메모",
]


def _entry(sample: Any, split: str, source: str, question_types: dict, forced_reason: str = "") -> dict:
    qrow = question_types.get(sample.id, {})
    return {
        "sample": sample, "split": split, "source": source, "forced_reason": forced_reason,
        "auto_type": qrow.get("auto_type", "UNKNOWN"), "matched_rule": qrow.get("matched_rule", ""),
    }


def select(cfg: Any, dev_samples: list, train_samples: list, question_types: dict,
           target_n: int = TARGET_N) -> list[dict]:
    from src.data import subset

    baseline = subset(list(dev_samples), bls.BASELINE_N, cfg.seed, "type")
    baseline_ids = {s.id for s in baseline}

    rows: dict[str, dict] = {}
    for s in baseline:
        rows[s.id] = _entry(s, "dev", "baseline300", question_types)

    remaining = [s for s in dev_samples if s.id not in baseline_ids]
    for r in bls._pick_supplement(remaining, question_types, baseline_ids, cfg.seed):
        rows.setdefault(r["sample"].id, _entry(r["sample"], "dev", "supplement", question_types))

    for r in bls._build_forced_rows(train_samples, dev_samples, question_types):
        sid = r["sample"].id
        if sid in rows:
            prev = rows[sid]["forced_reason"]
            rows[sid]["forced_reason"] = f"{prev}+{r['forced_reason']}" if prev else r["forced_reason"]
        else:
            rows[sid] = _entry(r["sample"], r["split"], "forced", question_types, r["forced_reason"])

    need = target_n - len(rows)
    if need > 0:
        def key(s: Any) -> str:
            return hashlib.sha256(f"{cfg.seed}:auditfill:{s.id}".encode()).hexdigest()

        pool = sorted((s for s in train_samples if s.id not in rows), key=key)
        for s in pool[:need]:
            rows[s.id] = _entry(s, "train", "train_random", question_types)
    elif need < 0:
        print(f"경고: baseline·보충·강제 포함만으로 {len(rows)}행이라 목표 {target_n}행을 넘었습니다. "
              f"자르지 않고 그대로 배정합니다.")

    selected = list(rows.values())
    bls._forbid_disallowed_split(selected, "라벨 감사 500 후보 전체")
    # split 표기와 별개로 id·경로에서도 test를 막는다(방어를 이중으로).
    leaked = [r["sample"].id for r in selected
              if r["sample"].id.lower().startswith("test") or "test" in r["sample"].image_path.parts]
    if leaked:
        raise bls.ForbiddenSplitError(
            f"test로 보이는 샘플 {len(leaked)}개가 섞였습니다 — Kaggle 4-b 위반. 예: {leaked[:5]}")
    return selected


def assign(selected: list[dict], annotators: list[str] = ANNOTATORS) -> list[dict]:
    """출처별로 id 순 정렬 후 라운드로빈 — 세 사람이 같은 개수, 같은 출처 비율을 받는다."""
    counter = 0
    for source in SOURCE_ORDER:
        for row in sorted((r for r in selected if r["source"] == source), key=lambda r: r["sample"].id):
            row["담당"] = annotators[counter % len(annotators)]
            counter += 1
    order = {name: i for i, name in enumerate(annotators)}
    return sorted(selected, key=lambda r: (order[r["담당"]], r["sample"].id))


def write_csv(rows: list[dict], path: Path = CSV_PATH) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:  # BOM: 엑셀·구글시트에서 한글 안 깨짐
        writer = csv.DictWriter(f, fieldnames=FIELDS)
        writer.writeheader()
        for i, row in enumerate(rows, start=1):
            s = row["sample"]
            choices = list(s.choices) + [""] * (4 - len(s.choices))
            writer.writerow({
                "번호": i, "row_id": s.id, "split": row["split"], "image_file": s.image_path.name,
                "질문": s.question,
                "선택지1": choices[0], "선택지2": choices[1], "선택지3": choices[2], "선택지4": choices[3],
                "제공된정답": s.answer or "",
                "auto_type": row["auto_type"], "matched_rule": row["matched_rule"],
                "source": row["source"], "forced_reason": row["forced_reason"],
                "담당": row["담당"],
                "유형확정": row["auto_type"],  # 맞으면 그대로 두면 된다
                "라벨정확": "", "판독가능": "", "메모": "",
            })


def write_xlsx(rows: list[dict], path: Path = XLSX_PATH) -> None:
    """작성용 엑셀. 첫 시트 '라벨링'에 드롭다운·필터·고정 행을 미리 걸어둔다.

    구글 드라이브에 올려 구글시트로 열면 드롭다운(데이터 확인란)이 그대로 유지된다. 첫 시트가
    '라벨링'이어야 시트를 CSV로 내려받을 때 이 시트가 나온다(merge_label_sheets.py 입력).
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "라벨링"
    ws.append(FIELDS)
    editable = {"유형확정", "라벨정확", "판독가능", "메모"}
    for col, h in enumerate(FIELDS, start=1):
        c = ws.cell(row=1, column=col)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="FFF2CC" if h in editable else "DCE6F1")

    for i, row in enumerate(rows, start=1):
        s = row["sample"]
        choices = list(s.choices) + [""] * (4 - len(s.choices))
        ws.append([i, s.id, row["split"], s.image_path.name, s.question, *choices, s.answer or "",
                   row["auto_type"], row["matched_rule"], row["source"], row["forced_reason"],
                   row["담당"], row["auto_type"], "", "", ""])

    last = len(rows) + 1
    col_of = {h: get_column_letter(i) for i, h in enumerate(FIELDS, start=1)}
    for header, options in [("유형확정", bls.TYPES), ("라벨정확", bls.LABEL_ACCURACY_OPTIONS),
                            ("판독가능", bls.READABILITY_OPTIONS)]:
        dv = DataValidation(type="list", formula1=f'"{",".join(options)}"', allow_blank=True,
                            showErrorMessage=True, errorTitle="허용값이 아닙니다",
                            error=f"드롭다운에서 고르세요: {', '.join(options)}")
        ws.add_data_validation(dv)
        dv.add(f"{col_of[header]}2:{col_of[header]}{last}")

    ws.freeze_panes = "F2"  # 헤더와 앞 5열(번호~질문)을 고정
    ws.auto_filter.ref = f"A1:{get_column_letter(len(FIELDS))}{last}"
    widths = {"번호": 6, "row_id": 12, "split": 7, "image_file": 16, "질문": 48,
              "선택지1": 20, "선택지2": 20, "선택지3": 20, "선택지4": 20, "제공된정답": 10,
              "auto_type": 12, "matched_rule": 22, "source": 13, "forced_reason": 14,
              "담당": 8, "유형확정": 12, "라벨정확": 10, "판독가능": 10, "메모": 30}
    for h, w in widths.items():
        ws.column_dimensions[col_of[h]].width = w
    for r in ws.iter_rows(min_row=2, max_row=last, min_col=5, max_col=5):
        r[0].alignment = Alignment(wrap_text=True, vertical="top")

    guide = wb.create_sheet("작업지침")
    lines = [
        ("라벨 감사 작업지침", True),
        ("", False),
        ("이 시트에는 train·dev만 있습니다. test 데이터는 없고 있어서도 안 됩니다 (Kaggle 기본규칙 4-b).", True),
        ("", False),
        ("1. 필터에서 '담당'이 자기 이름인 행만 봅니다. 다른 사람 행은 건드리지 않습니다.", False),
        ("2. 이미지 zip(label_audit_images_500.zip)을 풀고, 'image_file'과 같은 이름의 파일을 열어 확인합니다.", False),
        ("3. 노란 열 네 개만 채웁니다:", False),
        ("   유형확정 — 자동 분류(auto_type)가 맞으면 그대로 두고, 틀리면 드롭다운에서 고칩니다.", False),
        ("   라벨정확 — 제공된 정답이 이미지·질문과 맞는가: 맞음 / 틀림 / 애매. 이 열이 핵심입니다.", False),
        ("   판독가능 — 이미지에서 답을 읽어낼 수 있는가: 선명 / 흐림 / 판독불가.", False),
        ("   메모 — 애매하면 이유를 한 줄로.", False),
        ("4. 판단이 안 서면 억지로 확정하지 말고 '애매'로 두고 메모를 남깁니다. 우리가 재려는 게 그 비율입니다.", False),
        ("5. 작은 글씨는 원본을 확대해서 봅니다. 다 채우면 MM에 알립니다.", False),
        ("", False),
        ("source 열: baseline300/train_random은 무작위 표본, supplement/forced는 일부러 모은 의심·희소 표본입니다.", False),
        ("모두 똑같이 검수하면 됩니다. 집계에서 알아서 구분합니다.", False),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        c = guide.cell(row=i, column=1, value=text)
        c.font = Font(bold=bold, size=13 if i == 1 else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    guide.column_dimensions["A"].width = 110

    types_ws = wb.create_sheet("유형정의")
    types_ws.append(["유형", "정의", "예시1", "예시2"])
    for c in types_ws[1]:
        c.font = Font(bold=True)
    for name in bls.TYPES:
        definition, examples = bls.TYPE_DEFINITIONS[name]
        types_ws.append([name, definition, *(examples + ["", ""])[:2]])
    for col, w in zip("ABCD", [14, 45, 40, 40]):
        types_ws.column_dimensions[col].width = w

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_zip(rows: list[dict], password: str | None, path: Path = ZIP_PATH) -> int:
    """감사 대상 이미지만 평평한 zip으로 묶는다. 파일명(dev_0001.jpg 등)이 CSV의 image_file과 같다."""
    path.parent.mkdir(parents=True, exist_ok=True)
    missing = [r["sample"].image_path for r in rows if not r["sample"].image_path.exists()]
    if missing:
        raise SystemExit(f"이미지 {len(missing)}장을 찾을 수 없습니다. 예: {missing[:3]}")

    if password:
        try:
            import pyzipper
        except ImportError:
            raise SystemExit("암호 zip에는 pyzipper가 필요합니다: pip install pyzipper "
                             "(암호 없이 만들려면 --no-password)") from None
        zf = pyzipper.AESZipFile(path, "w", compression=pyzipper.ZIP_STORED, encryption=pyzipper.WZ_AES)
        zf.setpassword(password.encode("utf-8"))
    else:
        zf = zipfile.ZipFile(path, "w", compression=zipfile.ZIP_STORED)  # JPEG는 재압축 이득이 없다

    with zf:
        for r in rows:
            zf.write(r["sample"].image_path, arcname=r["sample"].image_path.name)
    return len(rows)


def summarize(rows: list[dict]) -> None:
    print(f"\n총 {len(rows)}행")
    print("  source :", dict(Counter(r["source"] for r in rows)))
    print("  split  :", dict(Counter(r["split"] for r in rows)))
    print("  담당   :", dict(Counter(r["담당"] for r in rows)))
    for name in ANNOTATORS:
        mine = [r for r in rows if r["담당"] == name]
        print(f"    {name}: {len(mine)}개, source={dict(Counter(r['source'] for r in mine))}")
    print("  auto_type:", dict(Counter(r["auto_type"] for r in rows).most_common()))


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--no-password", action="store_true", help="zip을 암호 없이 만든다")
    parser.add_argument("--target", type=int, default=TARGET_N)
    args = parser.parse_args()

    from src.config import load_config
    from src.data import load_split

    cfg = load_config(bls.CONFIG_NAME)
    dev_samples, _ = load_split(cfg, "dev")
    train_samples, _ = load_split(cfg, "train")
    question_types = bls._load_question_types()

    selected = select(cfg, list(dev_samples), list(train_samples), question_types, args.target)
    rows = assign(selected)

    password = None
    if not args.no_password:
        password = getpass.getpass("감사용 zip 암호를 정하세요 (저장소·시트에 적지 말고 MM DM으로만 공유): ")
        if not password:
            raise SystemExit("암호가 비어 있습니다. 암호 없이 만들려면 --no-password 를 쓰세요.")

    write_csv(rows)
    write_xlsx(rows)
    n_img = write_zip(rows, password)
    print(f"저장됨: {XLSX_PATH}")
    print(f"저장됨: {CSV_PATH}")
    print(f"저장됨: {ZIP_PATH} ({n_img}장, {'암호 있음' if password else '암호 없음'})")
    summarize(rows)
    print("\n다음: label_audit_500.xlsx를 팀 구글 드라이브에 올려 구글시트로 열고 공유한다(드롭다운 유지됨, "
          "첫 시트가 '라벨링'). 이미지 zip은 파일로 직접 전달한다 — git에 올리지 않는다.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
