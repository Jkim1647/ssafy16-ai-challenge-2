# -*- coding: utf-8 -*-
"""제출 여러 개가 서로 다른 test 문항만 모아 Excel 로 비교한다. GPU 미사용.

사용: python tools/compare_subs_xlsx.py T14 T16 T18 T19 T20 T21
시트: 요약(Public·H408 v2·서로 다른 답 수 행렬) / 다른문항(하나라도 다른 문항 — 질문·보기, 제출별 답과 그 답 확률, 구성 모델 답과 확신)
규칙 4-b: test 문항은 정답을 판단하거나 손으로 고쳐 제출하지 않는다. 제출끼리 어디서 갈리는지 보는 분석용이다.
"""
import csv, io, contextlib, os, sys, collections
import numpy as np
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
with contextlib.redirect_stdout(io.StringIO()):
    import export_test_models_xlsx_v2 as V  # noqa: E402  (T14~T20 결합 함수, Q35·Q38·9B 예측)
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

R, E, L, pr = V.R, V.E, 'abcd', V.E.pr
# 이름: (제출 파일, Public 또는 None, 결합 함수, 설명)
ALL = {
    'T6': ('T6_4pmean_test.csv', 0.97438, E.t6, 'R3·35B·397B·Gemma 확률 평균'),
    'T14': ('T14_t6_gemmarot_test.csv', 0.97378, R.SUBS['T14'][2], 'T6 에서 Gemma 를 보기 회전 4개 평균으로'),
    'T15': ('T15_t6_q38half_test.csv', 0.97467, R.t15, 'T6 + Qwen3.8-27B x0.5'),
    'T16': ('T16_t14half_w_test.csv', 0.97378, R.t16, '0.5·T14 + 397B·Gemma 1/6 + 35B·R3 1/12'),
    'T17': ('T17_t6_9bdora05_test.csv', 0.97408, V.SUBS['T17'][2], 'T6 + 9B DoRA x0.5'),
    'T18': ('T18_t6_q35half_test.csv', 0.97497, V.SUBS['T18'][2], 'T6 + Qwen3.5-27B LoRA 셔플 x0.5'),
    'T19': ('T19_r3toq35_test.csv', 0.97319, V.SUBS['T19'][2], 'T6 에서 R3 -> Qwen3.5-27B LoRA 셔플'),
    'T20': ('T20_t14_q35half_test.csv', 0.97438, V.SUBS['T20'][2], 'T14 + Qwen3.5-27B LoRA 셔플 x0.5'),
}
ALL['T21'] = ('T21_t16t18_19to1_test.csv', None, lambda D, i: (19 * R.t16(D, i) + V.SUBS['T18'][2](D, i)) / 20, '(19·T16 + T18)/20 확률 혼합')
S = sys.argv[1:] or ['T14', 'T16', 'T18', 'T19', 'T20', 'T21']
A = {k: {r['id']: r['answer'] for r in csv.DictReader(open('submissions/' + ALL[k][0], encoding='utf-8-sig'))} for k in S}
P = {k: {i: ALL[k][2](E.TEST, i) for i in E.tids} for k in S}
for k in S:
    assert all(L[int(np.argmax(P[k][i]))] == A[k][i] for i in E.tids), f'{k} 결합 함수가 제출 파일과 다르다'
MEM = [('R3', 'R3'), ('35B', '35B'), ('397B', '397B'), ('Gemma', 'Gemma'), ('Q35', 'Q35 (9번)'), ('Q38', 'Q38')]
diff = [i for i in E.tids if len({A[k][i] for k in S}) > 1]
diff.sort(key=lambda i: (-len({A[k][i] for k in S}), i))

wb = Workbook()
bold, wrap = Font(bold=True), Alignment(wrap_text=True, vertical='top')
fills = {x: PatternFill('solid', fgColor=c) for x, c in zip(L, ['FCE4D6', 'DDEBF7', 'E2EFDA', 'FFF2CC'])}
ws = wb.active
ws.title = '요약'
ws.append(['제출', '설명', 'Public', 'H408 v2(401)'] + [f'{b}와 다른 답' for b in S])
for c in ws[1]:
    c.font = bold
for a in S:
    h = sum(L[int(np.argmax(ALL[a][2](E.M, i)))] == E.g[i] for i in E.ids)
    ws.append([a, ALL[a][3], ALL[a][1] if ALL[a][1] else '미확인', h] + [sum(A[a][i] != A[b][i] for i in E.tids) for b in S])
ws.append([])
ws.append([f'하나라도 답이 다른 문항 {len(diff)}개 — 다른문항 시트. 보기 글자마다 색이 다르다(a 주황 · b 파랑 · c 초록 · d 노랑).'])
ws.append(['규칙 4-b: test 문항은 정답을 판단하거나 손으로 고쳐 제출하지 않는다. 분석용.'])
for j, w in enumerate([8, 44, 10, 12] + [11] * len(S), 1):
    ws.column_dimensions[get_column_letter(j)].width = w

ws = wb.create_sheet('다른문항')
head = ['id', '유형', '질문', 'a', 'b', 'c', 'd', '갈린 답 조합']
for k in S:
    head += [f'{k} 답', f'{k} 확률']
for _, lab in MEM:
    head += [f'{lab} 답', f'{lab} 확신']
ws.append(head)
for c in ws[1]:
    c.font = bold
    c.alignment = wrap
for i in diff:
    r = V.rows[i]
    grp = collections.defaultdict(list)
    for k in S:
        grp[A[k][i]].append(k)
    combo = ' / '.join(f"{x}: {','.join(v)}" for x, v in sorted(grp.items()))
    line = [i, V.T.get(i, '?'), r['question'], r['a'], r['b'], r['c'], r['d'], combo]
    for k in S:
        line += [A[k][i], round(float(P[k][i].max()), 3)]
    for m, _ in MEM:
        p = pr(E.TEST, m, i)
        line += [L[int(np.argmax(p))], round(float(p.max()), 3)]
    ws.append(line)
    rr = ws.max_row
    for s, k in enumerate(S):
        ws.cell(row=rr, column=9 + 2 * s).fill = fills[A[k][i]]
    for m_, (m, _) in enumerate(MEM):
        c = ws.cell(row=rr, column=9 + 2 * len(S) + 2 * m_)
        c.fill = fills[c.value]
    for j in (3, 4, 5, 6, 7, 8):
        ws.cell(row=rr, column=j).alignment = wrap
ws.freeze_panes = 'D2'
for j, w in enumerate([15, 11, 44, 16, 16, 16, 16, 30] + [7, 8] * len(S) + [7, 8] * len(MEM), 1):
    ws.column_dimensions[get_column_letter(j)].width = w
ws.auto_filter.ref = ws.dimensions
out = f"reports/compare_{'_'.join(S)}_20260925.xlsx"
wb.save(out)
print(out, '다른 문항', len(diff))
