# -*- coding: utf-8 -*-
"""test 6,714건 예측이 모두 있는 모델 13개를 T14 와 비교해 Excel 로 정리한다. GPU 미사용(저장된 확률만 쓴다).

시트
  설명        읽는 법과 규칙 4-b 주의
  모델요약    모델별 H408 v2·val667·단독 Public, T14 와 같은 답/다른 답 수, 평균 확신도
  다른문항    13개 중 하나라도 T14 와 답이 다른 문항 — 질문·보기, T14 답과 확신, 모델별 (답, 그 답 확률, T14 답 확률)
  전체문항    6,714건 전부 — 질문·보기, T14·T16 답과 a~d 확률, 모델 13개의 답과 a~d 확률(T14와 다른 답은 노란 칸)

확신도 = 그 모델의 a~d 로그확률을 softmax 한 뒤 고른 보기의 확률.
규칙 4-b: test 문항은 정답을 판단하거나 손으로 고치지 않는다. 이 표는 모델끼리 어디서 갈리는지 보는 용도다.
"""
import csv, io, contextlib, os, sys, collections
import numpy as np
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
with contextlib.redirect_stdout(io.StringIO()):
    import report_all_submissions as R  # noqa: E402  (T14 결합 함수, H408 v2 구성원)
from ensemble_submit import load_any  # noqa: E402
from ensemble_methods import lsm  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

E = R.E
L = 'abcd'
P = 'shared_predictions/'
# 이름: (파일, 설명, H408 v2 키 또는 None, val667, 단독 Public)
MODELS = {
    '397B': (P + 'nebius_397b_test/Qwen3.5-397B-A17B-FP8_read_test_tiles2x2x2_vllm_predictions.jsonl', 'Qwen3.5-397B-A17B 제로샷 판독+타일', '397B', 643, 0.96693),
    '35B_r2v3': (P + 'runpod_test_r2v3_20260924/Qwen3.6-35B-A3B_read_testTEST_r2v3_tiles2x2x2_r2v3_vllm_predictions.jsonl', 'Qwen3.6-35B-A3B 제로샷 + 라벨 규칙(r2v3)', '35Br2v3', None, None),
    '35B_r1v3': (P + 'runpod_e_20260924/Qwen3.6-35B-A3B_read_testTEST_r1v3_tiles2x2x2_r1v3_vllm_predictions.jsonl', 'Qwen3.6-35B-A3B 제로샷 + 구조화 판독(r1v3)', '35Br1v3', None, None),
    '35B': (P + 'colab_ai2_35b_out/Qwen3.6-35B-A3B_read_test_tiles2x2x2_predictions.jsonl', 'Qwen3.6-35B-A3B 제로샷 판독+타일', '35B', 644, 0.96455),
    'Gemma': (P + 'runpod_gemma4_31b/gemma_h408_test/test_predictions.jsonl', 'Gemma4-31B LoRA 타일', 'Gemma', 647, None),
    'R3': (P + 'runpod_27b_R3/R3_h408/test_predictions.jsonl', 'Qwen3.6-27B LoRA 타일(R3)', 'R3', 646, None),
    'Q38': (P + 'runpod_q38_20260924/Qwen3.8-27B_read_testTEST_tiles2x2x2_vllm_predictions.jsonl', 'Qwen3.8-27B 제로샷 판독+타일', 'Q38', 636, None),
    '27B_LoRA_full': (P + 'colab_ai2_35b_out/lora/27b_lora_r8/test_predictions.jsonl', 'Qwen3.6-27B LoRA 전체 이미지', None, 646, 0.96187),
    '9B_LoRA': (P + 'colab_ai2_35b_out/lora/Qwen3.5-9B_lora_r8_lr0.0001_ep1/test_predictions.jsonl', 'Qwen3.5-9B LoRA 전체 이미지', None, None, None),
    '27B_zs': (P + 'colab_ai2_35b_out/Qwen3.6-27B_ours_test_predictions.jsonl', 'Qwen3.6-27B 제로샷(우리 프롬프트, 전체 이미지)', None, 637, None),
    '35B_zs_ours': (P + 'colab_ai2_35b_out/Qwen3.6-35B-A3B_ours_test_predictions.jsonl', 'Qwen3.6-35B-A3B 제로샷(우리 프롬프트)', None, 630, None),
    '35B_zs_sw': (P + '20260921/test_35b_seokwoong.jsonl', 'Qwen3.6-35B-A3B 제로샷(석웅 프롬프트)', None, None, 0.95144),
    '9B_zs': (P + '20260921/test_9b_ours.jsonl', 'Qwen3.5-9B NF4 제로샷', None, None, 0.94221),
}
H408_OTHER = {'27B_LoRA_full': 341, '9B_LoRA': 321}  # Colab 채점값(reports/h408_lora7_20260924.md)

tids = E.tids
PR = {}
for n, (f, *_rest) in MODELS.items():
    d = load_any(f)
    assert all(i in d for i in tids), f'{n}: test 누락'
    PR[n] = {i: np.exp(lsm(d[i]['logprobs'])) for i in tids}
names = list(MODELS)
t14fn = R.SUBS['T14'][2]
t14p = {i: t14fn(E.TEST, i) for i in tids}
t14a = {i: R.TP['T14'][i] for i in tids}
assert all(L[int(np.argmax(t14p[i]))] == t14a[i] for i in tids), 'T14 결합 함수가 제출 파일과 다르다'
t16p = {i: R.SUBS['T16'][2](E.TEST, i) for i in tids}
t16a = {i: R.TP['T16'][i] for i in tids}
assert all(L[int(np.argmax(t16p[i]))] == t16a[i] for i in tids), 'T16 결합 함수가 제출 파일과 다르다'
ans = {n: {i: L[int(np.argmax(PR[n][i]))] for i in tids} for n in names}
rows = R.rows
T = E.T
g, ids = E.g, E.ids

wb = Workbook()
bold = Font(bold=True)
yellow = PatternFill('solid', fgColor='FFF2CC')
green = PatternFill('solid', fgColor='E2EFDA')
wrap = Alignment(wrap_text=True, vertical='top')

ws = wb.active
ws.title = '설명'
for line in [
    'test 예측이 6,714건 모두 있는 모델 13개와 T14(제출)를 비교한 표입니다.',
    '확신도 = 그 모델의 a~d 확률(로그확률을 softmax) 중 고른 보기의 확률. 1에 가까울수록 확신.',
    '전체문항 시트: test 6,714건 전부. 모델마다 [답, a, b, c, d 확률] 5칸, T14와 다른 답은 노란 칸.',
    '다른문항 시트: 13개 중 하나라도 T14 와 다른 답을 고른 문항. 노란 칸 = T14 와 다른 답, 초록 칸 = T14 와 같은 답.',
    '"T14답 확률" = 그 모델이 T14 가 고른 보기에 준 확률. 낮을수록 그 모델은 T14 답에 반대한다.',
    'T14 = T6(R3·35B·397B·Gemma 확률 평균)에서 Gemma 를 보기 회전 4개 평균으로 바꾼 것(Public 0.97378).',
    'T16 = 0.5·T14 + 397B·Gemma 각 1/6 + 35B·R3 각 1/12 확률 평균(Public 0.97378). T14 와 다른 T16 답은 노란 칸.',
    '규칙 4-b: test 문항은 정답을 판단하거나 손으로 고쳐 제출하지 않는다. 모델이 어디서 갈리는지 보는 분석용이다.',
    '만든 도구: tools/export_test_models_xlsx.py (저장된 예측만 사용, GPU 없음).']:
    ws.append([line])
ws.column_dimensions['A'].width = 120

ws = wb.create_sheet('모델요약')
ws.append(['모델', '설명', 'H408 v2(401)', 'val667', '단독 Public', 'T14와 같은 답', 'T14와 다른 답', '평균 확신도(전체)',
           '평균 확신도(T14와 다른 문항)', 'T14답에 준 평균 확률(다른 문항)'])
for c in ws[1]:
    c.font = bold
for n in names:
    f, desc, hk, v667, pub = MODELS[n]
    h = sum(L[int(np.argmax(E.pr(E.M, hk, i)))] == g[i] for i in ids) if hk else H408_OTHER.get(n)
    diff = [i for i in tids if ans[n][i] != t14a[i]]
    conf = [PR[n][i].max() for i in tids]
    ws.append([n, desc, h, v667, pub, len(tids) - len(diff), len(diff), round(float(np.mean(conf)), 4),
               round(float(np.mean([PR[n][i].max() for i in diff])), 4) if diff else None,
               round(float(np.mean([PR[n][i][L.index(t14a[i])] for i in diff])), 4) if diff else None])
for j, w in enumerate([16, 44, 12, 9, 11, 13, 13, 16, 22, 24], 1):
    ws.column_dimensions[get_column_letter(j)].width = w

ws = wb.create_sheet('다른문항')
diff_ids = [i for i in tids if t16a[i] != t14a[i] or any(ans[n][i] != t14a[i] for n in names)]
diff_ids.sort(key=lambda i: (sum(ans[n][i] == t14a[i] for n in names), i))
head = ['id', '유형', '질문', 'a', 'b', 'c', 'd', 'T14 답', 'T14 확신', 'T16 답', 'T16 확신', 'T14와 같은 모델 수(/13)', '가장 많이 나온 다른 답']
for n in names:
    head += [f'{n} 답', f'{n} 확신', f'{n} T14답 확률']
ws.append(head)
for c in ws[1]:
    c.font = bold
    c.alignment = wrap
for i in diff_ids:
    r = rows[i]
    same = sum(ans[n][i] == t14a[i] for n in names)
    other = collections.Counter(ans[n][i] for n in names if ans[n][i] != t14a[i]).most_common(1)
    line = [i, T.get(i, '?'), r['question'], r['a'], r['b'], r['c'], r['d'], t14a[i],
            round(float(t14p[i].max()), 3), t16a[i], round(float(t16p[i].max()), 3), same,
            f'{other[0][0]} ({other[0][1]}개)' if other else '']
    for n in names:
        line += [ans[n][i], round(float(PR[n][i].max()), 3), round(float(PR[n][i][L.index(t14a[i])]), 3)]
    ws.append(line)
    rr = ws.max_row
    for k, n in enumerate(names):
        cell = ws.cell(row=rr, column=len(head) - 3 * (len(names) - k) + 1)
        cell.fill = yellow if ans[n][i] != t14a[i] else green
    if t16a[i] != t14a[i]:
        ws.cell(row=rr, column=10).fill = yellow
ws.freeze_panes = 'D2'
for j, w in enumerate([15, 11, 50, 18, 18, 18, 18, 8, 9, 8, 9, 12, 14] + [8, 8, 10] * len(names), 1):
    ws.column_dimensions[get_column_letter(j)].width = w
ws.auto_filter.ref = ws.dimensions

ws = wb.create_sheet('전체문항')
head = ['id', '유형', '질문', 'a', 'b', 'c', 'd', 'T14 답', 'T14 a', 'T14 b', 'T14 c', 'T14 d',
        'T16 답', 'T16 a', 'T16 b', 'T16 c', 'T16 d', 'T14와 같은 모델 수(/13)']
for n in names:
    head += [f'{n} 답'] + [f'{n} {x}' for x in L]
ws.append(head)
for c in ws[1]:
    c.font = bold
    c.alignment = wrap
for i in tids:
    r = rows[i]
    line = [i, T.get(i, '?'), r['question'], r['a'], r['b'], r['c'], r['d'], t14a[i]]         + [round(float(v), 3) for v in t14p[i]] + [t16a[i]] + [round(float(v), 3) for v in t16p[i]]         + [sum(ans[n][i] == t14a[i] for n in names)]
    for n in names:
        line += [ans[n][i]] + [round(float(v), 3) for v in PR[n][i]]
    ws.append(line)
    rr = ws.max_row
    for k, n in enumerate(names):
        if ans[n][i] != t14a[i]:
            ws.cell(row=rr, column=19 + 5 * k).fill = yellow
    if t16a[i] != t14a[i]:
        ws.cell(row=rr, column=13).fill = yellow
ws.freeze_panes = 'D2'
for j, w in enumerate([15, 11, 50, 18, 18, 18, 18, 8, 7, 7, 7, 7, 8, 7, 7, 7, 7, 12] + [7, 6, 6, 6, 6] * len(names), 1):
    ws.column_dimensions[get_column_letter(j)].width = w
ws.auto_filter.ref = ws.dimensions

out = 'reports/test_13models_vs_T14_20260924.xlsx'
wb.save(out)
print(f'{out}: 다른문항 {len(diff_ids)} / 6714')
for n in names:
    print(f'  {n:14s} T14와 다른 답 {sum(ans[n][i] != t14a[i] for i in tids)}')
